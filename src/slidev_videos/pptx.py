"""`slidev-videos pptx`: a PowerPoint file from the stage recorder's takes.

    slidev-videos pptx <recording-dir> [--deck deck.md] [--out talk.pptx]
                       [--profile pptx|small|copy] [--cap 6M] [--encoder auto|nvenc|x264]
                       [--size 1920x1080] [--clips DIR] [--allow-skips] [--keep-encodes]
                       [--autoplay after|with] [--json]

One 16:9 slide per segment of `slidev-stage-record`'s index.json, in order
(slide, then click): the take full-bleed, its first frame as the poster, set
to start on its own when the slide is shown and to hold its last frame; no
media controls in the slide show; the slide's speaker notes from the deck
(a click segment's under a `[click K/N]` line: the part after Slidev's
`[click]` markers, or the whole notes again). Nothing on the slide is
editable: it is the recorded talk, a slide at a time.

Each take is re-encoded on the `pptx` profile — 4:2:0 H.264 High, NVENC VBR at
CQ 24 or x264 at CRF 21, both under `-maxrate 6M -bufsize 12M`, a key frame a
second, faststart — since the takes themselves are near lossless (30-60
Mbit/s); `small` is the plates' 4 Mbit/s, `copy` keeps the takes as recorded.
The takes are silent: a clip slide gets the source clip's sound, laid where
the recorder logged the clip playing (`clip.sound` in index.json), cut from
the clip in the deck's build (or `--clips`, or `public/videos/` beside the
deck) at the web tier's level (-16 LUFS), AAC, no loudnorm. The world's own
sounds (the hum, a Soundtrack) are not in the takes, and not in the PPTX.

A clip slide the recorder left to its clip (`slidev-stage-record
--skip-clips`: an `embed` segment, no take) gets the clip file itself, from
the deck's build or `public/videos/` (or `--clips`), else downloaded from its
URL (the release): as it is where PowerPoint plays it so (H.264, 4:2:0, 16:9,
within the cap, AAC), else encoded on the profile (cover-cropped to 16:9, as
the player fills the frame); its sound as the player played it (`volume`,
`muted`), the slide as long as the clip. A clip slide's ClipEnd line, exit push
or AfterFlight is not in it (nor in a filmed take, which ends on the clip).

A segment the recorder skipped (a clip from another origin, or one the
browser could not decode) stops the build, unless --allow-skips: then a black
slide names the clip. Exit 0 on a file written, 1 on a skip or an ffmpeg
failure, 2 on a usage error (no index.json, no python-pptx).

Needs the `pptx` extra (python-pptx): `pip install "slidev-videos[pptx]"`.
"""
from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

from . import deck as _deck

W, H = 12192000, 6858000                 # 16:9 in EMU (13.333 in × 7.5 in)
SHOW_MEDIA_CTRLS = "{2FDB2607-1784-4EEB-B798-7EB5836EED8A}"
P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"

# The encoder options of each profile, by codec, and the rate cap; mirrors
# packages/stage/bin/lib/record-ffmpeg.mjs (`small` is its plate profile).
PROFILES = {
    "pptx": {
        "h264_nvenc": ["-preset", "p5", "-tune", "hq", "-rc", "vbr", "-cq", "24", "-b:v", "0"],
        "libx264": ["-preset", "medium", "-crf", "21"],
        "cap": "6M",
    },
    "small": {
        "h264_nvenc": ["-preset", "p5", "-tune", "hq", "-rc", "vbr", "-cq", "27", "-b:v", "0"],
        "libx264": ["-preset", "medium", "-crf", "23"],
        "cap": "4M",
    },
    "copy": None,
}
INSTALL = 'pip install "slidev-videos[pptx]"  (python-pptx)'


def _num(x: float) -> str:
    """A number for an ffmpeg filter: at most 3 decimals, no trailing zeros."""
    s = f"{float(x):.3f}".rstrip("0").rstrip(".")
    return s if s not in ("", "-0") else "0"


def _double(rate: str) -> str:
    m = re.fullmatch(r"(\d+(?:\.\d+)?)([kKmM]?)", rate)
    if not m:
        raise ValueError(f"--cap wants a rate such as 6M or 4500k (got {rate})")
    return f"{_num(float(m.group(1)) * 2)}{m.group(2)}"


def cap_bps(rate: str) -> int:
    """A rate such as 6M or 4500k in bits a second."""
    m = re.fullmatch(r"(\d+(?:\.\d+)?)([kKmM]?)", rate)
    if not m:
        raise ValueError(f"--cap wants a rate such as 6M or 4500k (got {rate})")
    return round(float(m.group(1)) * {"": 1, "k": 1e3, "m": 1e6}[m.group(2).lower()])


def clip_plan(info: dict, *, profile: str, cap: str | None, size: tuple[int, int] | None,
              muted: bool, volume: float) -> dict:
    """How an embedded clip goes into the PPTX, each stream: `copy` where
    PowerPoint plays it as it is (H.264 4:2:0, 16:9, within the cap, the size
    asked for; any clip under --profile copy), else `encode`; the sound `none`
    (muted, silent, or no sound), `volume` (the player's quieter level),
    `copy` (AAC) or `encode`. `why` names what made the video an encode."""
    why = []
    if PROFILES[profile] is not None:
        w, h, br = info.get("width"), info.get("height"), info.get("bit_rate")
        if info.get("vcodec") != "h264":
            why.append(f"{info.get('vcodec')}, not H.264")
        if info.get("pix_fmt") not in ("yuv420p", "yuvj420p"):
            why.append(f"{info.get('pix_fmt')}, not 4:2:0")
        if br and cap and br > cap_bps(cap):
            why.append(f"{br / 1e6:.1f} Mbit/s, over the cap {cap}")
        if w and h and abs(w / h - 16 / 9) > 0.01:
            why.append(f"{w}x{h}, not 16:9")
        if size and (w, h) != tuple(size):
            why.append(f"{w}x{h}, not {size[0]}x{size[1]}")
    if muted or not (volume or 0) > 0 or not info.get("acodec"):
        audio = "none"
    elif volume < 1:
        audio = "volume"
    else:
        audio = "copy" if info.get("acodec") == "aac" else "encode"
    return {"video": "encode" if why else "copy", "audio": audio, "why": why}


def _cover(info: dict, size: tuple[int, int] | None) -> tuple[int, int] | None:
    """The frame an encoded clip is cover-cropped to: --size, else 16:9 at the
    clip's width when it is not 16:9; None to keep it."""
    if size:
        return size
    w, h = info.get("width"), info.get("height")
    if w and h and abs(w / h - 16 / 9) > 0.01:
        return w, round(w * 9 / 16 / 2) * 2
    return None


def embed_args(ffmpeg: str, clip: Path, out: Path, plan: dict, *, codec: str, profile: str, info: dict,
               cap: str | None = None, size: tuple[int, int] | None = None, volume: float = 1) -> list[str]:
    """The ffmpeg command line that puts a clip slide's clip into the PPTX (clip_plan's plan)."""
    args = [ffmpeg, "-hide_banner", "-loglevel", "error", "-y", "-i", str(clip), "-map", "0:v:0"]
    prof = PROFILES[profile]
    if plan["video"] == "copy" or prof is None:
        args += ["-c:v", "copy"]
    else:
        rate = cap or prof["cap"]
        args += ["-c:v", codec, *prof[codec], "-maxrate", rate, "-bufsize", _double(rate),
                 "-profile:v", "high", "-pix_fmt", "yuv420p", "-g", str(round(info.get("fps") or 25))]
        target = _cover(info, size)
        if target:
            tw, th = target
            args += ["-vf", f"scale={tw}:{th}:force_original_aspect_ratio=increase,crop={tw}:{th}"]
    if plan["audio"] == "none":
        args += ["-an"]
    elif plan["audio"] == "copy":
        args += ["-map", "0:a:0", "-c:a", "copy"]
    else:
        args += ["-map", "0:a:0", *(["-af", f"volume={_num(volume)}"] if plan["audio"] == "volume" else []),
                 "-c:a", "aac", "-b:a", "192k"]
    return args + ["-movflags", "+faststart", str(out)]


def probe_clip(ffprobe: str, f: Path) -> dict:
    """What clip_plan reads of a clip: its video codec, pixel format, size, rate
    (the video stream's, else the file's), sound codec, length, frame rate."""
    r = subprocess.run([ffprobe, "-v", "error", "-show_entries",
                        "stream=codec_type,codec_name,pix_fmt,width,height,bit_rate,avg_frame_rate:format=duration,bit_rate",
                        "-of", "json", str(f)], capture_output=True, text=True)
    try:
        data = json.loads(r.stdout or "{}")
    except ValueError:
        data = {}
    streams = data.get("streams") or []
    v = next((x for x in streams if x.get("codec_type") == "video"), {})
    a = next((x for x in streams if x.get("codec_type") == "audio"), {})
    fmt = data.get("format") or {}

    def num(x):
        try:
            return float(x)
        except (TypeError, ValueError):
            return None
    n, _, d = str(v.get("avg_frame_rate") or "").partition("/")
    fps = num(n) / num(d) if num(n) and num(d) else None
    br = num(v.get("bit_rate")) or num(fmt.get("bit_rate"))
    return {"vcodec": v.get("codec_name"), "pix_fmt": v.get("pix_fmt"), "width": v.get("width"), "height": v.get("height"),
            "bit_rate": round(br) if br else None, "acodec": a.get("codec_name"), "duration": num(fmt.get("duration")), "fps": fps}


def fetch_clip(url: str, into: Path) -> tuple[Path | None, str | None]:
    """A clip not on disk, downloaded from its URL (the talk's release) → (the file, None) or (None, why)."""
    import urllib.request
    into.mkdir(parents=True, exist_ok=True)
    out = into / (Path(url.split("?")[0].split("#")[0]).name or "clip.mp4")
    part = out.with_name(out.name + ".part")
    try:
        req = urllib.request.Request(url, headers={"User-Agent": "slidev-videos"})
        with urllib.request.urlopen(req, timeout=60) as r, open(part, "wb") as f:
            shutil.copyfileobj(r, f, 1 << 20)
    except Exception as e:  # noqa: BLE001 — any failure is the same: no clip
        part.unlink(missing_ok=True)
        return None, str(e)
    os.replace(part, out)
    return out, None


def audible(sound: dict | None) -> bool:
    return bool(sound) and not sound.get("muted") and (sound.get("volume") or 0) > 0


def audio_filter(sound: dict, seconds: float, fps: float) -> str:
    """The clip's sound from its own second `t`, for as long as it played
    (`end` - `start` and one frame), placed at the take's second `start`,
    padded with silence and cut to the take."""
    dur = sound["end"] - sound["start"] + 1 / fps
    parts = [f"atrim=start={_num(sound.get('t') or 0)}:duration={_num(dur)}", "asetpts=PTS-STARTPTS"]
    if sound["start"] > 0:
        parts.append(f"adelay={round(sound['start'] * 1000)}:all=1")
    vol = sound.get("volume", 1)
    if vol < 1:
        parts.append(f"volume={_num(vol)}")
    parts += ["apad", f"atrim=end={_num(seconds)}"]
    return "[1:a:0]" + ",".join(parts) + "[a]"


def encode_args(ffmpeg: str, take: Path, out: Path, *, codec: str, profile: str, fps: float, seconds: float,
                cap: str | None = None, size: tuple[int, int] | None = None,
                sound: dict | None = None, clip: Path | None = None) -> list[str]:
    """The ffmpeg command line that turns one take into the PPTX's video."""
    args = [ffmpeg, "-hide_banner", "-loglevel", "error", "-y", "-i", str(take)]
    with_sound = audible(sound) and clip is not None
    if with_sound:
        if sound.get("loop"):
            args += ["-stream_loop", "-1"]
        args += ["-i", str(clip)]
    args += ["-map", "0:v:0"]
    prof = PROFILES[profile]
    if prof is None:
        args += ["-c:v", "copy"]
    else:
        rate = cap or prof["cap"]
        args += ["-c:v", codec, *prof[codec], "-maxrate", rate, "-bufsize", _double(rate),
                 "-profile:v", "high", "-pix_fmt", "yuv420p", "-g", str(round(fps))]
        if size:
            args += ["-vf", f"scale={size[0]}:{size[1]}:flags=lanczos"]
    if with_sound:
        args += ["-filter_complex", audio_filter(sound, seconds, fps), "-map", "[a]", "-c:a", "aac", "-b:a", "192k"]
    else:
        args += ["-an"]
    return args + ["-movflags", "+faststart", str(out)]


def find_clip(src: str, dist: Path | None, dirs: list[Path]) -> Path | None:
    """The source clip of a recorded clip slide: by its path in the deck's
    build, else by its name in each of `dirs`."""
    rel = re.sub(r"^[a-z]+://[^/]+/", "", src).split("?")[0].split("#")[0].lstrip("/")
    if dist and rel and (dist / rel).is_file():
        return dist / rel
    for d in dirs:
        if (d / Path(rel).name).is_file():
            return d / Path(rel).name
    return None


def has_audio(ffprobe: str, f: Path) -> bool:
    r = subprocess.run([ffprobe, "-v", "error", "-select_streams", "a", "-show_entries", "stream=index",
                        "-of", "csv=p=0", str(f)], capture_output=True, text=True)
    return r.returncode == 0 and bool(r.stdout.strip())


def autoplay_timing(pic_id: int, ms: int, start: str = "after", vol: int = 100000) -> str:
    """`p:timing` for a video that starts on its own when the slide is shown
    ("Start: Automatically"): the main sequence's first effect is a
    `mediacall` playFrom(0.0) on the picture, begun with the sequence; the
    `p:video` node keeps fill="hold", so the last frame stays (no rewind)."""
    node = {"after": "afterEffect", "with": "withEffect"}[start]
    tgt = f'<p:tgtEl><p:spTgt spid="{pic_id}"/></p:tgtEl>'
    sld = '<p:tgtEl><p:sldTgt/></p:tgtEl>'
    return (
        f'<p:timing xmlns:p="{NS_P}"><p:tnLst><p:par>'
        '<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
        '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
        '<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/>'
        '<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
        '<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
        f'<p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="{node}">'
        '<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
        f'<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="{ms}" fill="hold"/>{tgt}</p:cBhvr></p:cmd>'
        '</p:childTnLst></p:cTn></p:par>'
        '</p:childTnLst></p:cTn></p:par>'
        '</p:childTnLst></p:cTn></p:par>'
        '</p:childTnLst></p:cTn>'
        f'<p:prevCondLst><p:cond evt="onPrev" delay="0">{sld}</p:cond></p:prevCondLst>'
        f'<p:nextCondLst><p:cond evt="onNext" delay="0">{sld}</p:cond></p:nextCondLst></p:seq>'
        f'<p:video><p:cMediaNode vol="{vol}"><p:cTn id="7" fill="hold" display="0">'
        f'<p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn>{tgt}</p:cMediaNode></p:video>'
        '</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'
    )


def _hide_media_controls(prs) -> None:
    """presProps: p:showPr with p14:showMediaCtrls val="0" (the slide show's
    play bar under a video, off)."""
    from lxml import etree
    part = next(p for p in prs.part.package.iter_parts() if str(p.partname) == "/ppt/presProps.xml")
    root = etree.fromstring(part.blob)
    p = lambda t: f"{{{NS_P}}}{t}"  # noqa: E731
    show = root.find(p("showPr"))
    if show is None:
        show = etree.Element(p("showPr"), showNarration="1")
        etree.SubElement(show, p("present"))
        etree.SubElement(show, p("sldAll"))
        pen = etree.SubElement(show, p("penClr"))
        etree.SubElement(pen, f"{{{NS_A}}}prstClr", val="red")
        after = [c for c in root if c.tag in (p("clrMru"), p("extLst"))]
        (after[0].addprevious(show) if after else root.append(show))
    ext_lst = show.find(p("extLst"))
    if ext_lst is None:
        ext_lst = etree.SubElement(show, p("extLst"))
    for e in ext_lst.findall(p("ext")):
        if e.get("uri") == SHOW_MEDIA_CTRLS:
            ext_lst.remove(e)
    ext = etree.SubElement(ext_lst, p("ext"), uri=SHOW_MEDIA_CTRLS)
    etree.SubElement(ext, f"{{{P14}}}showMediaCtrls", nsmap={"p14": P14}, val="0")
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


_XML_BAD = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _black(slide) -> None:
    from pptx.dml.color import RGBColor
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(0, 0, 0)


def _skipped_slide(slide, seg: dict) -> None:
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Pt
    box = slide.shapes.add_textbox(Emu(W // 12), Emu(H // 3), Emu(W * 10 // 12), Emu(H // 3))
    tf = box.text_frame
    tf.word_wrap = True
    tf.text = f"Slide {seg['slide']}: the clip was not recorded\n{(seg.get('clip') or {}).get('src', '')}\n{seg.get('note', '')}"
    for para in tf.paragraphs:
        for run in para.runs:
            run.font.size = Pt(20)
            run.font.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)


def _run(args: list[str]) -> str | None:
    r = subprocess.run(args, capture_output=True, text=True)
    return None if r.returncode == 0 else (r.stderr.strip().splitlines() or [f"exit {r.returncode}"])[-1]


def _pick_codec(ffmpeg: str, want: str, profile: str) -> str:
    if PROFILES[profile] is None:
        return "copy"
    from .pipeline import _encodes_with
    if want != "x264" and _encodes_with(ffmpeg, "h264_nvenc", *PROFILES[profile]["h264_nvenc"]):
        return "h264_nvenc"
    if want == "nvenc":
        raise RuntimeError("h264_nvenc does not work with this ffmpeg")
    return "libx264"


def _size(v: str | None) -> tuple[int, int] | None:
    if not v:
        return None
    m = re.fullmatch(r"(\d+)x(\d+)", v)
    if not m:
        raise argparse.ArgumentTypeError(f"--size wants WxH, e.g. 1920x1080 (got {v})")
    return int(m.group(1)), int(m.group(2))


def _label(s: dict) -> str:
    return f"slide {s['slide']}" + (f" click {s['clicks']}" if s.get("clicks") else "")


def _default_deck() -> Path | None:
    from . import pipeline
    root = pipeline._PROJECT.root if pipeline._PROJECT else None
    if root is None:
        return None
    if (root / "deck.md").is_file():
        return root / "deck.md"
    mds = sorted(root.glob("*.md"))
    return mds[0] if len(mds) == 1 else None


def cmd_pptx(args: argparse.Namespace) -> int:
    from . import pipeline
    report = pipeline._report(args)
    rec = Path(args.recording)
    index = rec / "index.json"
    if not index.is_file():
        print(f"error: no index.json in {rec} (a slidev-stage-record out-dir)", file=sys.stderr)
        return 2
    try:
        import pptx  # noqa: F401
        from pptx import Presentation
        from pptx.oxml import parse_xml
        from pptx.oxml.ns import qn
    except ImportError as e:
        print(f"error: slidev-videos pptx needs python-pptx ({e}): {INSTALL}", file=sys.stderr)
        return 2
    try:
        size = _size(args.size)
        if args.cap:
            _double(args.cap)
    except (ValueError, argparse.ArgumentTypeError) as e:
        print(f"error: {e}", file=sys.stderr)
        return 2
    data = json.loads(index.read_text(encoding="utf-8"))
    segs = sorted(data.get("segments") or [], key=lambda s: (s["slide"], s.get("clicks") or 0))
    out = Path(args.out) if args.out else rec.with_name(rec.name + ".pptx")
    deck_path = Path(args.deck) if args.deck else _default_deck()
    skipped = [s for s in segs if s.get("skipped")]
    report.update(out=str(out), profile=args.profile, skipped=[
        {"slide": s["slide"], "clicks": s.get("clicks") or 0, "clip": (s.get("clip") or {}).get("src")} for s in skipped])
    if not segs:
        print(f"error: {index} lists no segments", file=sys.stderr)
        return 1
    if skipped and not args.allow_skips:
        for s in skipped:
            print(f"error: slide {s['slide']} was not recorded: {s.get('note', 'skipped')} ({(s.get('clip') or {}).get('src')})",
                  file=sys.stderr)
        print("  pull the clips first (slidev-videos pull --include-shared), build local-first "
              "(VITE_VIDEOS_LOCAL_FIRST=1) and record again; or --allow-skips for a black slide naming the clip",
              file=sys.stderr)
        return 1
    if deck_path and not deck_path.is_file():
        print(f"error: no deck {deck_path}", file=sys.stderr)
        return 2
    notes = _deck.notes_by_slide(deck_path) if deck_path else {}
    if not deck_path:
        print("  ! no deck (--deck): the slides go without notes", file=sys.stderr)
    clicks = {}
    for s in segs:
        clicks[s["slide"]] = max(clicks.get(s["slide"], 0), s.get("clicks") or 0)

    ffmpeg, ffprobe = pipeline._ffmpeg(), pipeline._ffprobe()
    try:
        codec = _pick_codec(ffmpeg, args.encoder, args.profile)
    except RuntimeError as e:
        print(f"error: {e}", file=sys.stderr)
        return 1
    fps = float(data.get("fps") or 50)
    dist = Path(data["deck"]) if data.get("deck") else None
    dirs = [Path(d) for d in args.clips or []]
    if deck_path:
        dirs.append(deck_path.parent / "public" / "videos")
    cap = None if PROFILES[args.profile] is None else (args.cap or PROFILES[args.profile]["cap"])
    report.update(encoder=codec, cap=cap, sound=[], embedded=[], warnings=[])
    enc_dir = out.parent / (out.stem + ".enc")
    enc_dir.mkdir(parents=True, exist_ok=True)

    built = []           # (segment, encoded mp4, poster) or (segment, None, None) for a skip
    loud = set()         # the files that carry sound (the player's level in the slide show)
    print(f"{len(segs)} segment(s) from {rec} -> {out}  ({args.profile}, {codec}{f', cap {cap}' if cap else ''})")
    for s in segs:
        if s.get("skipped"):
            built.append((s, None, None))
            print(f"  {_label(s)}: skipped, a black slide naming {(s.get('clip') or {}).get('src')}")
            continue
        if s.get("embed"):
            got = _embed(s, args, ffmpeg, ffprobe, codec, cap, size, dist, dirs, enc_dir, report)
            if got is None:
                return 1
            built.append(got)
            if got[0]["audio"]:
                loud.add(got[0]["file"])
            continue
        take = rec / s["file"]
        if not take.is_file():
            print(f"error: {take} is missing", file=sys.stderr)
            return 1
        sound = (s.get("clip") or {}).get("sound")
        clip = None
        say = "silent"
        if audible(sound):
            src = s["clip"]["src"]
            clip = find_clip(src, dist, dirs)
            if clip is None:
                w = f"{s['file']}: the clip {src} is not in the build nor in {', '.join(map(str, dirs)) or '--clips'}: silent"
                report["warnings"].append(w)
                print(f"  ! {w}", file=sys.stderr)
            elif not has_audio(ffprobe, clip):
                clip, say = None, f"silent ({Path(src).name} has no sound)"
            else:
                say = f"with {Path(src).name}'s sound from {sound['start']:.2f} s"
                loud.add(s["file"])
                report["sound"].append({"file": s["file"], "clip": src, "start": sound["start"]})
        enc = enc_dir / s["file"]
        poster = enc.with_suffix(".jpg")
        argv = encode_args(ffmpeg, take, enc, codec=codec, profile=args.profile, fps=fps, seconds=s["seconds"],
                           cap=args.cap, size=size, sound=sound, clip=clip)
        sig = json.dumps([argv, take.stat().st_mtime_ns, clip.stat().st_mtime_ns if clip else None])
        stamp = enc.with_suffix(".args")
        if not (enc.is_file() and poster.is_file() and stamp.is_file() and stamp.read_text() == sig):
            err = _run(argv) or _run([ffmpeg, "-hide_banner", "-loglevel", "error", "-y", "-i", str(enc),
                                      "-frames:v", "1", "-q:v", "3", str(poster)])
            if err:
                print(f"error: ffmpeg on {take.name}: {err}", file=sys.stderr)
                return 1
            stamp.write_text(sig)
        else:
            say += ", encoded before"
        built.append((s, enc, poster))
        print(f"  {s['file']:<12} {_label(s)}  {s['seconds']} s  {pipeline.human_size(enc.stat().st_size)}  {say}")

    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]
    for s, enc, poster in built:
        slide = prs.slides.add_slide(blank)
        _black(slide)
        if enc is None:
            _skipped_slide(slide, s)
        else:
            pic = slide.shapes.add_movie(str(enc), 0, 0, W, H, poster_frame_image=str(poster), mime_type="video/mp4")
            sld = slide._element
            old = sld.find(qn("p:timing"))
            vol = 100000 if s["file"] in loud else 80000
            new = parse_xml(autoplay_timing(pic.shape_id, round(s["seconds"] * 1000), args.autoplay, vol))
            old.addprevious(new)
            sld.remove(old)
        k = s.get("clicks") or 0
        text = _deck.notes_for(notes, s["slide"], k, clicks.get(s["slide"], 0))
        if text:
            slide.notes_slide.notes_text_frame.text = _XML_BAD.sub("", text)
    _hide_media_controls(prs)
    part = out.with_name(out.name + ".part")
    prs.save(str(part))
    os.replace(part, out)
    if not args.keep_encodes:
        shutil.rmtree(enc_dir, ignore_errors=True)
    seconds = round(sum(s["seconds"] for s, enc, _ in built if enc is not None), 3)
    report.update(slides=len(built), segments=len(segs), bytes=out.stat().st_size, seconds=seconds)
    print(f"{out}: {len(built)} slide(s), {seconds} s of video, {pipeline.human_size(out.stat().st_size)}"
          f"{f', {len(skipped)} skipped' if skipped else ''}; open it in PowerPoint and check the first slide plays on its own")
    return 0


def _embed(s: dict, args, ffmpeg: str, ffprobe: str, codec: str, cap: str | None, size, dist: Path | None,
           dirs: list[Path], enc_dir: Path, report: dict) -> tuple[dict, Path, Path] | None:
    """A clip slide left to its clip: the clip found (or downloaded), put in as
    it is or encoded (clip_plan), its poster → (the segment as built, the mp4,
    the poster), or None on a failure (said)."""
    c = s.get("clip") or {}
    src = c.get("src") or ""
    name = f"{s['slide']:02d}" + (f"-c{s['clicks']}" if s.get("clicks") else "") + ".mp4"
    clip, origin, why = find_clip(src, dist, dirs), "local", None
    if clip is None and re.match(r"https?://", src):
        clip, why = fetch_clip(src, enc_dir / "clips")
        origin = "url"
    if clip is None:
        print(f"error: {_label(s)}: the clip {src} is not in the build nor in {', '.join(map(str, dirs)) or '--clips'}"
              f"{f', and its URL failed ({why})' if why else ''}; pull the clips (slidev-videos pull --include-shared)",
              file=sys.stderr)
        return None
    info = probe_clip(ffprobe, clip)
    if not info.get("duration") or not info.get("vcodec"):
        print(f"error: {_label(s)}: ffprobe reads no video in {clip}", file=sys.stderr)
        return None
    volume = c.get("volume") if isinstance(c.get("volume"), (int, float)) else 1
    plan = clip_plan(info, profile=args.profile, cap=cap, size=size, muted=bool(c.get("muted")), volume=volume)
    enc = enc_dir / name
    poster = enc.with_suffix(".jpg")
    argv = embed_args(ffmpeg, clip, enc, plan, codec=codec, profile=args.profile, info=info, cap=cap, size=size, volume=volume)
    sig = json.dumps([argv, clip.stat().st_mtime_ns, clip.stat().st_size])
    stamp = enc.with_suffix(".args")
    say = ""
    if not (enc.is_file() and poster.is_file() and stamp.is_file() and stamp.read_text() == sig):
        err = _run(argv) or _run([ffmpeg, "-hide_banner", "-loglevel", "error", "-y", "-i", str(enc),
                                  "-frames:v", "1", "-q:v", "3", str(poster)])
        if err:
            print(f"error: ffmpeg on {clip.name}: {err}", file=sys.stderr)
            return None
        stamp.write_text(sig)
    else:
        say = ", encoded before"
    seconds = round(probe_clip(ffprobe, enc).get("duration") or info["duration"], 3)
    done = {**s, "file": name, "seconds": seconds, "audio": plan["audio"] != "none"}
    report["embedded"].append({"slide": s["slide"], "clicks": s.get("clicks") or 0, "clip": src, "from": origin,
                               "video": plan["video"], "audio": plan["audio"], "seconds": seconds,
                               **({"why": plan["why"]} if plan["why"] else {})})
    print(f"  {name:<12} {_label(s)}  {seconds} s  {pipeline_size(enc)}  the clip {Path(src).name} ({origin}): "
          f"video {plan['video']}{' (' + '; '.join(plan['why']) + ')' if plan['why'] else ''}, sound {plan['audio']}{say}")
    return done, enc, poster


def pipeline_size(f: Path) -> str:
    from . import pipeline
    return pipeline.human_size(f.stat().st_size)


def add_parser(sub, parents=()) -> None:
    p = sub.add_parser("pptx", parents=list(parents),
                       help="a PPTX from slidev-stage-record's takes: one autoplaying full-bleed video per slide and click, the deck's notes")
    p.add_argument("recording", help="the recorder's out-dir (NN.mp4, NN-cK.mp4, index.json)")
    p.add_argument("--deck", default=None, help="deck.md for the speaker notes (default: the project's deck.md)")
    p.add_argument("--out", default=None, help="the .pptx to write (default: <recording>.pptx beside the dir)")
    p.add_argument("--profile", choices=list(PROFILES), default="pptx",
                   help="pptx: CQ 24 / CRF 21 under 6 Mbit/s (default); small: CQ 27 / CRF 23 under 4 Mbit/s; copy: the takes as recorded")
    p.add_argument("--cap", default=None, help="the rate cap, e.g. 6M (default: the profile's); the buffer is twice it")
    p.add_argument("--encoder", choices=("auto", "nvenc", "x264"), default="auto", help="auto: NVENC where it encodes, else x264")
    p.add_argument("--size", default=None, help="scale every take to WxH (default: as recorded)")
    p.add_argument("--clips", action="append", metavar="DIR", help="where the source clips are, when not in the deck's build")
    p.add_argument("--allow-skips", action="store_true", help="a skipped segment becomes a black slide naming its clip")
    p.add_argument("--keep-encodes", action="store_true", help="keep <out>.enc/ (re-encodes, posters) and reuse it next time")
    p.add_argument("--autoplay", choices=("after", "with"), default="after",
                   help="the start effect's node: after (PowerPoint's Automatically; default) or with (With Previous)")
    p.set_defaults(func=cmd_pptx)
